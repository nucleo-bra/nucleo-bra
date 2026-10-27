# Converte o "Modelo Padrão Contrato BRA Ambientes.docx" (campos em dourado
# entre colchetes) no modelo do Núcleo (public/modelos/contrato-bra.docx),
# trocando cada campo pelo {campo} que o Núcleo preenche.
# Dados pessoais (CPF/RG/endereço) viram campos vindos de Contratos › Configurações,
# para não ficarem num arquivo público.
#
# Uso: python3 scripts/converter-modelo-bra.py "<modelo original>.docx"
import sys
from copy import deepcopy
from docx import Document
from docx.shared import Pt, RGBColor

ORIGEM = sys.argv[1]
SAIDA = __file__.rsplit("/", 1)[0] + "/../public/modelos/contrato-bra.docx"
d = Document(ORIGEM)
P = d.paragraphs


def texto_runs(p, *textos):
    """Define o texto dos runs do parágrafo na ordem; runs extras ficam vazios."""
    for i, r in enumerate(p.runs):
        r.text = textos[i] if i < len(textos) else ""


def achar(trecho, inicio=0):
    for i in range(inicio, len(P)):
        if trecho in P[i].text:
            return P[i]
    raise SystemExit(f"não achei: {trecho}")


def trocar_run(p, antigo, novo):
    for r in p.runs:
        if antigo in r.text:
            r.text = r.text.replace(antigo, novo)
            return
    raise SystemExit(f"não achei o run {antigo!r} em {p.text[:60]!r}")


def ancora(p, txt):
    r = p.add_run(txt)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r.font.size = Pt(6)


# capa
aviso = achar("MODELO PADRÃO — USO INTERNO")
texto_runs(aviso, "{#unidade}Contrato referente a unidade {unidade}{/unidade}")
instr = achar("Preencher os campos destacados")
instr._p.getparent().remove(instr._p)
trocar_run(achar("Nº do Contrato:"), "[Nº]", "{numero}")
trocar_run(achar("Cliente: [NOME DO CLIENTE]"), "[NOME DO CLIENTE]", "{cliente_nome}")

# partes
rep = achar("Representante legal:")
texto_runs(rep, rep.runs[0].text, "{rep_nome}, empresário, inscrito no CPF sob o nº {rep_cpf}, portador do RG nº {rep_rg}, "
           "residente e domiciliado na {rep_endereco}.")
trocar_run(achar("Nome completo:"), "[NOME DO CLIENTE]", "{cliente_nome}")
trocar_run(achar("CPF/CNPJ: [CPF/CNPJ DO CLIENTE]"), "[CPF/CNPJ DO CLIENTE]", "{cliente_documento}")
trocar_run(achar("Data de nascimento:"), "[DD/MM/AAAA]", "{cliente_nascimento}")
trocar_run(achar("Endereço: [ENDEREÇO"), "[ENDEREÇO COMPLETO DO CLIENTE]", "{cliente_endereco}")

# pagamento
trocar_run(achar("O preço total de venda"), "R$ [VALOR TOTAL POR EXTENSO]", "{valor_total} ({valor_extenso})")
trocar_run(achar("Forma de recebimento:"), "[BANCO / CHAVE PIX / TITULAR]", "{forma_recebimento}")
t = d.tables[0]
linha = t.rows[1].cells
texto_runs(linha[0].paragraphs[0], "{#condicoes}{condicao}")
texto_runs(linha[1].paragraphs[0], "{valor}")
texto_runs(linha[2].paragraphs[0], "{vencimento}{/condicoes}")
t._tbl.remove(t.rows[2]._tr)

# garantia, escopo, data
trocar_run(achar("Durante o prazo de garantia de"), "[12/24]", "{garantia_meses}")
# o escopo vai num parágrafo próprio, alinhado à esquerda (aceita várias linhas)
desc = achar("Descrever neste campo")
trocar_run(desc, "[DETALHAR ESCOPO ESPECÍFICO DESTA CONTRATAÇÃO]", "")
novo = deepcopy(desc._p)
desc._p.addnext(novo)
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH
esc = Paragraph(novo, desc._parent)
texto_runs(esc, "", "{escopo}")
esc.alignment = WD_ALIGN_PARAGRAPH.LEFT
desc.runs[0].text = desc.runs[0].text.rstrip()
trocar_run(achar("Condições específicas de fornecimento"), "[DESCREVER, SE HOUVER]", "{condicoes_especificas}")
data = achar("Paulínia, [DATA]")
trocar_run(data, "Paulínia, ", "{cidade}, ")
trocar_run(data, "[DATA]", "{data}")

# assinaturas (+ âncoras brancas para o Auto-Place do DocuSign)
i = P.index(data)
cli = achar("[NOME DO CLIENTE]", i)
texto_runs(cli, "{cliente_nome}")
texto_runs(achar("CPF: [CPF DO CLIENTE]", i), "CPF/CNPJ: {cliente_documento}")
fer = achar("Fernando", i)
texto_runs(fer, "{rep_nome}")
mat = achar("Matheus Esturrari", i)
texto_runs(mat, "{bra_nome}")
texto_runs(P[P.index(mat) + 1], "CPF: ", "{bra_cpf}")
ancora(P[P.index(cli) - 1], "/ass_cliente/")
ancora(P[P.index(fer) - 1], "/ass_rep/")
ancora(P[P.index(mat) - 1], "/ass_bra/")

# cabeçalho das páginas internas
for s in d.sections:
    for p in s.header.paragraphs:
        for r in p.runs:
            if "MODELO PADRÃO DE CONTRATO" in r.text:
                r.text = r.text.replace("MODELO PADRÃO DE CONTRATO", "Contrato {numero}")

d.core_properties.title = "Contrato de compra, venda e instalação de sistema — BRA Ambientes"
d.save(SAIDA)
resto = [p.text for p in d.paragraphs if "[" in p.text]
print("gerado:", SAIDA, "| colchetes restantes:", resto)
